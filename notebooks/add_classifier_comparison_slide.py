"""Add LogReg vs SetFit comparison slide to CA-3 presentation."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt

BASE = os.path.dirname(os.path.dirname(__file__))
PPT = os.path.join(BASE, "docs", "B_Tech_Project_CA3_Second_Presentation.pptx")
PPT_OUT = os.path.join(BASE, "docs", "B_Tech_Project_CA3_With_Classifier_Comparison.pptx")
OBJECT = 1

ROWS = [
    ("Aspect", "LogReg", "SetFit"),
    ("Approach", "Frozen MiniLM + linear classifier", "Contrastive fine-tuned MiniLM"),
    ("Val Accuracy", "73.15%", "79.87%"),
    ("Test Accuracy", "70.00%", "70.67%"),
    ("Holdout Accuracy", "67.23%", "62.18%"),
    ("Val→Test gap", "3.15 pts (stable)", "9.20 pts (overfits)"),
    ("Avg confidence", "0.53 (underconfident)", "0.86 (overconfident)"),
    ("Model size", "0.01 MB", "87.36 MB"),
    ("Training time", "~2–5 min", "~60–75 min"),
    ("Inference speed", "~12.6 ms/ticket", "~12.5 ms/ticket"),
    ("Best categories", "Database, Application, Storage", "Access, Infrastructure, Network, Security"),
    ("Recommendation", "Default for deployment", "Research; needs larger dataset"),
]


def main():
    prs = Presentation(PPT)
    slide = prs.slides.add_slide(prs.slide_layouts[OBJECT])
    slide.shapes.title.text = "LogReg vs SetFit — Classifier Comparison"

    rows, cols = len(ROWS), 3
    left, top, width, height = Inches(0.4), Inches(1.3), Inches(9.2), Inches(5.2)
    table = slide.shapes.add_table(rows, cols, left, top, width, height).table

    col_widths = [Inches(2.4), Inches(3.4), Inches(3.4)]
    for i, w in enumerate(col_widths):
        table.columns[i].width = w

    for r, row in enumerate(ROWS):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(11 if r > 0 else 12)
                p.font.name = "Times New Roman"
                if r == 0:
                    p.font.bold = True

    out = PPT_OUT
    try:
        prs.save(PPT)
        out = PPT
    except PermissionError:
        prs.save(PPT_OUT)
    print(f"Added comparison slide -> {out} ({len(prs.slides)} slides total)")


if __name__ == "__main__":
    main()
